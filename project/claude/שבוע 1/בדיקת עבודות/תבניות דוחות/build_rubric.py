# -*- coding: utf-8 -*-
"""המחוון לקריאה אנושית, ב-Word וב-Excel, באותו עיצוב כמו דוחות המשוב."""
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_reports as BR  # noqa: E402
import rubric_data as R  # noqa: E402

OUT = sys.argv[1]
FILL = {"s": "FFF2CC", "t": "E2EFDA", "b": "F2F2F2", "h": "DDEBF7"}
PART_HEAD = ["חלק", "נקודות", "סוג", "התשובה המלאה", "ניקוד חלקי והורדות"]
PART_W = [3.6, 1.5, 1.4, 9.8, 9.8]


# ---------------- Word ----------------
def table(doc, head, rows, widths, fills=None, bold_first=True):
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.style = "Table Grid"
    for j, h in enumerate(head):
        BR.cell_text(t.rows[0].cells[j], h, True, 10, FILL["h"], center=(len(h) < 8))
    for i, row in enumerate(rows, 1):
        f = fills[i - 1] if fills else None
        for j, v in enumerate(row):
            BR.cell_text(t.rows[i].cells[j], v or "", bold_first and j == 0, 9.5, f, center=(j in (1, 2) and len(head) == 5))
    BR.rtl_table(t, widths)
    BR.para(doc, "", space_after=4)


def parts_table(doc, rows):
    body, fills = [], []
    for name, pts, kind, full, partial, typ in rows:
        if typ == "h":
            body.append((name, "", "", "", ""))
            fills.append(FILL["h"])
        else:
            body.append((name, pts, kind, full, partial))
            fills.append(FILL.get(typ))
    table(doc, PART_HEAD, body, PART_W, fills)


def build_word(path):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    sec._sectPr.append(OxmlElement("w:bidi"))
    doc.styles["Normal"].font.name = BR.FONT
    doc.styles["Normal"].font.size = Pt(10.5)

    def h1(t):
        BR.para(doc, t, 16, True, "1F4E79", space_after=4)

    def h2(t):
        BR.para(doc, t, 12.5, True, "1F4E79", space_after=3)

    def brk():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    h1(R.TITLE)
    BR.para(doc, R.SUB, 9.5, color="555555", space_after=8)
    h2("סולם הניקוד")
    table(doc, ["אחוז", "מתי"], R.SCALE, [2.5, 23.6])
    h2("כללים קבועים")
    table(doc, ["כלל", "מה עושים"], R.RULES, [5.5, 20.6])
    h2("הורדות")
    table(doc, ["מקרה", "כמה יורד", "הערה"], R.DEDUCT, [7, 8, 11.1])
    h2("החלטות שהתקבלו בכיתה יב4")
    table(doc, ["משימה", "החלטה"], R.CLASS_DECISIONS, [4, 22.1])
    brk()
    h1("תרגילים 1-2: שיעור 1 ושיעור 2")
    BR.para(doc, "צהוב: צילום מסך. אפור: בונוס. ירוק: סך הכול.", 9.5, color="555555", space_after=4)
    parts_table(doc, R.L12)
    h2("מפתח תשובות, תרגילים 1-2")
    table(doc, ["נתון", "ערך"], R.KEY12, [8, 18.1])
    brk()
    h1("תרגיל 3: השוואה הוגנת ומאמץ דגימה")
    BR.para(doc, "צהוב: צילום מסך. אפור: בונוס. ירוק: סך הכול.", 9.5, color="555555", space_after=4)
    parts_table(doc, R.L3)
    h2("מפתח תשובות, תרגיל 3")
    table(doc, R.KEY3_HEAD, R.KEY3, [4.5, 5.4, 5.4, 5.4, 5.4])
    BR.para(doc, R.KEY3_NOTE, 9.5, color="555555")
    doc.save(path)


# ---------------- Excel ----------------
def xfill(hexc):
    return PatternFill("solid", fgColor=hexc)


def sheet(wb, title, head, rows, widths, fills=None, note=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    for i, (h, w) in enumerate(zip(head, widths), 1):
        c = ws.cell(1, i, h)
        c.font = Font(name=BR.FONT, bold=True, color="FFFFFF")
        c.fill = BR.H_FILL
        c.alignment = BR.WRAP
        c.border = BR.BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, 2):
        f = fills[r - 2] if fills else None
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.font = Font(name=BR.FONT, bold=(i == 1 or (f == FILL["h"])))
            c.alignment = BR.RIGHT if not (len(head) == 5 and i in (2, 3)) else Alignment(horizontal="center", vertical="top", wrap_text=True)
            c.border = BR.BOX
            if f:
                c.fill = xfill(f)
        longest = max(len(str(v or "")) for v in row)
        ws.row_dimensions[r].height = max(20, 15 * (1 + longest // 55))
    ws.freeze_panes = "B2"
    if note:
        r = len(rows) + 3
        c = ws.cell(r, 1, note)
        c.font = Font(name=BR.FONT, italic=True)
        c.alignment = BR.RIGHT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(head))
        ws.row_dimensions[r].height = 30
    return ws


def parts_rows(rows):
    body, fills = [], []
    for name, pts, kind, full, partial, typ in rows:
        body.append((name, "" if pts is None else pts, kind or "", full or "", partial or ""))
        fills.append(FILL["h"] if typ == "h" else FILL.get(typ))
    return body, fills


def build_excel(path):
    wb = Workbook()
    wb.remove(wb.active)
    pw = [24, 9, 9, 60, 60]
    legend = "צהוב: צילום מסך. אפור: בונוס. ירוק: סך הכול."
    b, f = parts_rows(R.L12)
    sheet(wb, "תרגילים 1-2", PART_HEAD, b, pw, f, legend)
    b, f = parts_rows(R.L3)
    sheet(wb, "תרגיל 3", PART_HEAD, b, pw, f, legend)
    rows = [("סולם", a, b) for a, b in R.SCALE] + [("כלל קבוע", a, b) for a, b in R.RULES]
    sheet(wb, "כללים", ["נושא", "שם", "מה עושים"], rows, [12, 26, 90])
    sheet(wb, "הורדות והחלטות", ["מקרה", "כמה יורד", "הערה"],
          R.DEDUCT + [("", "", "")] + [("החלטה בכיתה: " + a, "", b) for a, b in R.CLASS_DECISIONS], [40, 34, 70])
    sheet(wb, "מפתח 1-2", ["נתון", "ערך"], R.KEY12, [26, 50])
    sheet(wb, "מפתח 3", R.KEY3_HEAD, R.KEY3, [14, 18, 18, 18, 18], note=R.KEY3_NOTE)
    wb.save(path)


os.makedirs(OUT, exist_ok=True)
build_word(os.path.join(OUT, "מחוון לקריאה - PBI שבוע 1.docx"))
build_excel(os.path.join(OUT, "מחוון לקריאה - PBI שבוע 1.xlsx"))
print("ok")
