# -*- coding: utf-8 -*-
"""שיעור 3: משוב ב-Word לכל קבוצה, וגיליונות "שיעור 3" ו"סיכום הגשות" בקובצי ה-Excel של שיעורים 1-2."""
import os
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_reports as BR  # noqa: E402
import grading_data as G  # noqa: E402

OUT = sys.argv[1]
BR.OUT = OUT
TITLE3 = "תרגיל 3 PBI, יב4"
IDENT_PEN = 2
NOTEXT = {s["name"] for s in G.STUDENTS if s["summary"] == G.S_NOTEXT}

PARTS = [("L3.Q1", "שאלה 1: שתי השערות לירידה ב-2020", 6),
         ("L3.Q2a", "שאלה 2א: לכידות, סבבים וממוצע בסבב", 20),
         ("L3.Q2b", "שאלה 2ב: פחות חיות או מאמץ דגימה קטן יותר", 14),
         ("L3.Q2c", "שאלה 2ג: 2023 מול 2025", 8),
         ("L3.Q2d", "שאלה 2ד: טענה, ראיה ונימוק", 18),
         ("L3.C", "שלמות ההגשה", 4)]

Q2B = ("ההשוואה בין הממוצעים נכונה: 3 בסבב ב-2019 ו-2.5 ב-2020. ההבדל קטן, ולכן אפשר להסיק שהירידה מ-36 ל-5 "
       "לכידות נובעת בעיקר מכך שנערכו רק 2 סבבים. התשובה השאירה את המסקנה פתוחה (90%).")
Q2D = ("טענה: נכתב שמספר הלכידות ירד. זה נכון, אבל רק מתאר את הנתון. הטענה המצופה היא שב-2025 היו פחות "
       "חולדות קנגורו של אורד בשטח (3.5 מתוך 6). ראיה: מספרית ומדויקת (6 מתוך 6). נימוק: נכתב שאולי השתנה "
       "מאמץ הדגימה, אבל בשתי השנים נערכו 11 סבבים, ולכן הוא לא מסביר את הירידה (2 מתוך 6).")
Q2A = "כל הערכים נכונים. ב-2025 DO לא נלכד, ולכן הממוצע הוא 0 בסבב."
Q2C = "נכון: הממוצע ירד מ-13.6 בסבב ל-0."
C = "כל הסעיפים נענו, ויש טקסט ושני צילומים."
SG = "גרף מלא: Slicer עם DO, עמודה לכל שנה לפי הסדר, וספירה (Count) של recordID."
ST = "טבלת סבבים בעמוד נפרד, עם הערכים הנכונים: 12 ב-2019, 2 ב-2020, 11 ב-2023 וב-2025."

VERSIONS = {
 "A": dict(
   parts=[(6, "שתי השערות סבירות: הקורונה, ומחסור בכוח אדם או בציוד. שתיהן עוסקות במאמץ הדגימה. כדאי לשקול גם השערה מסוג אחר, למשל שהיו באמת פחות חיות."),
          (20, Q2A), (12.5, Q2B), (8, Q2C), (11.5, Q2D), (4, C)],
   sg=(15, SG), st=(15, ST),
   obs="גרף DO מסודר לפי שנה עם Slicer. טבלת סבבים מלאה (2017: 13). קבצים זהים אצל ארבעת התלמידים."),
 "B": dict(
   parts=[(6, "שתי השערות סבירות: הקורונה, וירידה בכוח האדם ובמשאבים. שתיהן עוסקות במאמץ הדגימה. כדאי לשקול גם השערה מסוג אחר, למשל שהיו באמת פחות חיות."),
          (20, Q2A), (12.5, Q2B), (8, Q2C), (11.5, Q2D), (4, C)],
   sg=(15, SG + " ב-Slicer מופיע (Blank), כלומר הקובץ לא נשען על הניקוי משיעור 2. לא ירדו על זה נקודות."),
   st=(15, ST + " ב-2017 מופיעים 14 סבבים במקום 13, כי הסבב עם period שלילי לא סונן. לא ירדו נקודות, כי השנים שנבדקו נכונות."),
   obs="גרף DO ועוד גרף DM מההדגמה. ב-Slicer מופיע (Blank) ובטבלה 2017 = 14: הניקוי משיעור 2 לא הוחל. קבצים זהים אצל תשעת התלמידים."),
}
SUMMARY = ("החישוב של ממוצע הלכידות בסבב נכון בכל ארבע השנים, והגרף וטבלת הסבבים בנויים כנדרש. הנקודות ירדו בעיקר "
           "בשאלה 2ד: בשתי השנים נערכו 11 סבבים, ולכן הירידה ב-2025 לא נובעת ממאמץ הדגימה, ואפשר לטעון שהיו פחות "
           "חיות מהמין בשטח. בשאלה 2ב הממוצעים כמעט שווים (3 ו-2.5), וזה מספיק כדי להסיק שהירידה במספר הלכידות "
           "ב-2020 נובעת בעיקר ממספר הסבבים. לשיפור: שני הסבבים של 2020 היו בחורף. כדאי לבדוק אם העונה משפיעה "
           "על מספר הלכידות של DO.")
DED_NOTE = ("שני קובצי הצילום, הגרף וטבלת הסבבים, זהים לקבצים שהגישו תלמידים אחרים: 2 נקודות פחות על כל אחד, "
            "יחד 4. כל תלמיד מצלם את המסך שלו.")

L3 = {"A": ["תלמיד לדוגמה 1"], "B": ["תלמיד לדוגמה 2"]}
TEACHER = {"A": "הערה למורה בלבד.", "B": "הערה למורה בלבד."}


def l3_student(name, g):
    v = VERSIONS[g]
    return dict(name=name, g=g, parts=v["parts"], sg=v["sg"], st=v["st"], identical=["graph", "table"],
                foreign=False, obs=v["obs"], teacher=TEACHER[g])


def scores3(s):
    text = sum(p for p, _ in s["parts"])
    scr = s["sg"][0] + s["st"][0]
    ded = IDENT_PEN * len(s["identical"]) + (1 if s["foreign"] else 0)
    return text, scr, ded, max(0, min(100, int(text + scr - ded + 0.5)))


def build_word3(g, studs, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    sec._sectPr.append(OxmlElement("w:bidi"))
    doc.styles["Normal"].font.name = BR.FONT
    doc.styles["Normal"].font.size = Pt(11)
    for n, s in enumerate(studs):
        text, scr, ded, final = scores3(s)
        BR.para(doc, "%s, %s" % (TITLE3, G.GROUPS[g]), 9, color="666666", space_after=0)
        BR.para(doc, "משוב: %s" % s["name"], 16, True, "1F4E79", space_after=6)
        BR.para(doc, "ציון סופי: %d מתוך 100" % final, 14, True, space_after=2)
        BR.para(doc, "חלק כתוב: %s מתוך 70. צילומי מסך: %s מתוך 30. הורדות: %s." % (BR.fmt(text), BR.fmt(scr), BR.fmt(ded)), 10.5, space_after=6)
        rows = [(lab, "%s מתוך %s" % (BR.fmt(p), mx), note) for (pid, lab, mx), (p, note) in zip(PARTS, s["parts"])]
        rows += [("צילום הגרף", "%s מתוך 15" % BR.fmt(s["sg"][0]), s["sg"][1]),
                 ("צילום טבלת הסבבים", "%s מתוך 15" % BR.fmt(s["st"][0]), s["st"][1]),
                 ("הורדות", "%s-" % BR.fmt(ded), DED_NOTE),
                 ("ציון סופי", "%d מתוך 100" % final, "חלק כתוב %s, צילומים %s, הורדות %s. מעוגל למספר שלם." % (BR.fmt(text), BR.fmt(scr), BR.fmt(ded)))]
        t = doc.add_table(rows=1 + len(rows), cols=3)
        t.style = "Table Grid"
        for j, h in enumerate(["חלק", "נקודות", "הסבר"]):
            BR.cell_text(t.rows[0].cells[j], h, True, 10, "DDEBF7", center=(j == 1))
        for i, (a, pts, note) in enumerate(rows, 1):
            fill = "FFF2CC" if a.startswith("צילום") else ("E2EFDA" if a == "ציון סופי" else None)
            BR.cell_text(t.rows[i].cells[0], a, a == "ציון סופי", 9, fill)
            BR.cell_text(t.rows[i].cells[1], pts, True, 9.5, fill, center=True)
            BR.cell_text(t.rows[i].cells[2], note, False, 9, fill)
        BR.rtl_table(t, [4.0, 2.5, 10.9])
        BR.para(doc, "", space_after=2)
        BR.para(doc, "סיכום", 12, True, "1F4E79", space_after=2)
        BR.para(doc, SUMMARY, 10.5, space_after=4)
        if n < len(studs) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)


def header(ws, row, labels, widths):
    for i, (h, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row, i, h)
        c.font = Font(name=BR.FONT, bold=True, color="FFFFFF")
        c.fill = BR.H_FILL
        c.alignment = BR.WRAP
        c.border = BR.BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 60


def style_row(ws, r, ncols, text_cols=(1,), fills=None):
    for i in range(1, ncols + 1):
        c = ws.cell(r, i)
        c.border = BR.BOX
        c.font = Font(name=BR.FONT, bold=(i == 1))
        c.alignment = BR.RIGHT if i in text_cols else Alignment(horizontal="center", vertical="center", wrap_text=True)
        if fills and i in fills:
            c.fill = fills[i]


def add_l3_sheets(path, g, studs3):
    wb = load_workbook(path)
    # ---- שיעור 3 ----
    ws = wb.create_sheet("שיעור 3", 1)
    ws.sheet_view.rightToLeft = True
    labels = ["שם התלמיד"] + ["%s (%s)" % (lab, mx) for _, lab, mx in PARTS] + [
        "סה\"כ כתוב (70)", "צילום הגרף (15)", "צילום טבלת הסבבים (15)", "סה\"כ צילומים (30)",
        "גרף זהה", "טבלה זהה", "שם קובץ של אחר", "סה\"כ הורדות", "בונוס (עד 5)", "ציון סופי (100)",
        "מה ראיתי בצילומים", "הערה למורה (לא לתלמיד)"]
    header(ws, 1, labels, [16] + [11] * 6 + [10, 10, 10, 10, 9, 9, 9, 9, 8, 11, 50, 50])
    dv = DataValidation(type="list", formula1='"כן,לא"', allow_blank=True)
    ws.add_data_validation(dv)
    L = get_column_letter
    for r, s in enumerate(studs3, 2):
        ws.cell(r, 1, s["name"])
        for j, (p, _) in enumerate(s["parts"]):
            ws.cell(r, 2 + j, p)
        ws.cell(r, 8, "=SUM(B%d:G%d)" % (r, r))
        ws.cell(r, 9, s["sg"][0])
        ws.cell(r, 10, s["st"][0])
        ws.cell(r, 11, "=I%d+J%d" % (r, r))
        ws.cell(r, 12, "כן" if "graph" in s["identical"] else "לא")
        ws.cell(r, 13, "כן" if "table" in s["identical"] else "לא")
        ws.cell(r, 14, "כן" if s["foreign"] else "לא")
        for col in (12, 13, 14):
            dv.add("%s%d" % (L(col), r))
        ws.cell(r, 15, '=%d*((L%d="כן")+(M%d="כן"))+(N%d="כן")' % (IDENT_PEN, r, r, r))
        ws.cell(r, 16, 0)
        ws.cell(r, 17, "=MAX(0,MIN(100,ROUND(H%d+K%d+P%d-O%d,0)))" % (r, r, r, r))
        ws.cell(r, 18, s["obs"])
        ws.cell(r, 19, s["teacher"])
        style_row(ws, r, 19, text_cols=(1, 18, 19),
                  fills={8: BR.TOT_FILL, 9: BR.SCR_FILL, 10: BR.SCR_FILL, 11: BR.TOT_FILL, 15: BR.TOT_FILL, 17: BR.TOT_FILL})
        ws.row_dimensions[r].height = 60
    ws.freeze_panes = "B2"
    note = ws.cell(len(studs3) + 3, 1, "כל \"כן\" בעמודות של גרף זהה ושל טבלה זהה מוריד 2 נקודות, ו\"כן\" בשם קובץ של אחר מוריד נקודה. "
                   "את כל הנקודות אפשר לשנות, והציון הסופי מתעדכן לבד.")
    note.font = Font(name=BR.FONT, italic=True)
    note.alignment = BR.RIGHT
    ws.merge_cells(start_row=len(studs3) + 3, start_column=1, end_row=len(studs3) + 3, end_column=17)

    # ---- סיכום הגשות ----
    w12 = wb["ציונים"]
    hdr = [c.value for c in w12[2]]
    c_final = hdr.index("ציון סופי (100)") + 1
    c_text = hdr.index("סה\"כ חלק כתוב (70)") + 1
    rows12 = {w12.cell(r, 1).value: r for r in range(3, w12.max_row + 1) if w12.cell(r, 1).value and w12.cell(r, c_final).value}
    rows3 = {s["name"]: r for r, s in enumerate(studs3, 2)}
    roster = sorted(set(rows12) | set(rows3))
    sm = wb.create_sheet("סיכום הגשות", 0)
    sm.sheet_view.rightToLeft = True
    labels = ["שם התלמיד", "ציון תרגילים 1-2", "ציון תרגיל 3", "סטטוס הגשה",
              "ממוצע לפי מה שהוגש", "ציון משוקלל (חסר = 0)", "הוגש טקסט בתרגילים 1-2"]
    header(sm, 1, labels, [18, 12, 12, 34, 14, 14, 12])
    q12, q3 = "'ציונים'!", "'שיעור 3'!"
    for r, name in enumerate(roster, 2):
        sm.cell(r, 1, name)
        sm.cell(r, 2, "=%s%s%d" % (q12, L(c_final), rows12[name]) if name in rows12 else "")
        sm.cell(r, 3, "=%sQ%d" % (q3, rows3[name]) if name in rows3 else "")
        sm.cell(r, 7, ("לא" if name in NOTEXT else "כן") if name in rows12 else "")
        sm.cell(r, 4, ('=IF(AND(B{r}<>"",C{r}<>""),"הגיש הכול",IF(B{r}<>"","חסר: תרגיל 3",IF(C{r}<>"","חסר: תרגילים 1-2","לא הגיש")))'
                       '&IF(G{r}="לא"," (בתרגילים 1-2 אין טקסט)","")').format(r=r))
        sm.cell(r, 5, '=IF(AND(B{r}="",C{r}=""),"",ROUND((2*N(B{r})+N(C{r}))/(IF(B{r}="",0,2)+IF(C{r}="",0,1)),1))'.format(r=r))
        sm.cell(r, 6, "=ROUND((2*N(B{r})+N(C{r}))/3,1)".format(r=r))
        style_row(sm, r, 7, text_cols=(1, 4), fills={5: BR.TOT_FILL, 6: BR.TOT_FILL})
    n = len(roster) + 3
    for i, line in enumerate([
        "איך מחושב: תרגילים 1-2 מכסים שני שיעורים, ותרגיל 3 שיעור אחד. לכן תרגילים 1-2 שווים פי 2.",
        "ממוצע לפי מה שהוגש: רק המשימות שהתלמיד הגיש נכנסות לחישוב. מתאים לראות את רמת העבודה.",
        "ציון משוקלל (חסר = 0): משימה שלא הוגשה נספרת 0. מתאים לציון תקופתי.",
        "הסטטוס מתעדכן לבד לפי הציונים בגיליונות \"ציונים\" ו\"שיעור 3\". תלמיד שלא הגיש אף משימה לא מופיע ברשימה: מוסיפים אותו בשורה חדשה."]):
        c = sm.cell(n + i, 1, line)
        c.font = Font(name=BR.FONT, italic=(i > 0), bold=(i == 0))
        c.alignment = BR.RIGHT
        sm.merge_cells(start_row=n + i, start_column=1, end_row=n + i, end_column=7)
        sm.row_dimensions[n + i].height = 30
    sm.freeze_panes = "B2"
    wb.active = 0
    wb.save(path)
    return roster


def main():
    BR.main()
    for g in ("A", "B"):
        studs3 = [l3_student(nm, g) for nm in L3[g]]
        build_word3(g, studs3, os.path.join(OUT, "משוב לתלמידים - %s, %s.docx" % (TITLE3, G.GROUPS[g])))
        name = "%s, %s" % (BR.TITLE, G.GROUPS[g])
        old = os.path.join(OUT, "ציונים - %s.xlsx" % name)
        new = os.path.join(OUT, "ציונים - תרגילים 1-3 PBI, יב4, %s.xlsx" % G.GROUPS[g])
        os.replace(old, new)
        roster = add_l3_sheets(new, g, studs3)
        for s in studs3:
            print(g, s["name"], scores3(s))
        print(g, "roster", len(roster))


main()
