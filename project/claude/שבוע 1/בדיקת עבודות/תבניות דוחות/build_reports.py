# -*- coding: utf-8 -*-
"""בונה לכל קבוצה קובץ Excel וקובץ Word מנתוני הבדיקה ב-grading_data.py."""
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import grading_data as G  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else ""
TITLE = "תרגילים 1-2 PBI, יב4"
FONT = "Arial"

L1S = [("הכרטיס מציג 15,120", 8), ("Count ולא Sum, לא מקוצר", 4)]
L2S = [("4 הצעדים הראשונים", 3), ("סינון period", 5), ("סינון species (Remove Empty)", 5),
       ("Choose Columns", 2.5), ("רשימה נקייה", 2.5)]


def fmt(x):
    return str(int(x)) if float(x).is_integer() else str(x)


def totals(st):
    pts = [p for p, _ in st["parts"]]
    l1, l2 = sum(pts[:6]), sum(pts[6:])
    b = st.get("bonus", (0, 0, ""))
    return l1, l2, b[0] + b[1]


def scr(st):
    (c1, n1), (c2, n2) = G.SCR[st["name"]]
    return c1, n1, c2, n2


def deductions(st):
    return G.IDENT_PENALTY * len(st["identical"]) + (1 if st["foreign"] else 0)


def final_grade(st):
    l1, l2, b = totals(st)
    c1, _, c2, _ = scr(st)
    raw = l1 + l2 + b + sum(c1) + sum(c2) - deductions(st)
    return max(0, min(100, int(raw + 0.5)))


def words(t):
    return len(t.split())


# ---------------- Excel ----------------
THIN = Side(style="thin", color="B7B7B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
H_FILL = PatternFill("solid", fgColor="1F4E79")
TXT_FILL = PatternFill("solid", fgColor="DDEBF7")
SCR_FILL = PatternFill("solid", fgColor="FFF2CC")
TOT_FILL = PatternFill("solid", fgColor="E2EFDA")
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="top", wrap_text=True)


def build_excel(group, studs, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ציונים"
    ws.sheet_view.rightToLeft = True

    cols = [("שם התלמיד", None, "name")]
    for pid, label, mx in G.L1_PARTS:
        cols.append((label.replace("שיעור 1, ", "").replace("שיעור 1: ", "") + " (%s)" % fmt(mx), mx, "t"))
    cols.append(("סה\"כ כתוב שיעור 1 (38)", 38, "l1"))
    for pid, label, mx in G.L2_PARTS:
        cols.append((label.replace("שיעור 2, ", "").replace("שיעור 2: ", "") + " (%s)" % fmt(mx), mx, "t"))
    cols.append(("סה\"כ כתוב שיעור 2 (32)", 32, "l2"))
    cols.append(("סה\"כ חלק כתוב (70)", 70, "text"))
    cols.append(("בונוס (עד 5)", 5, "bonus"))
    for lab, mx in L1S:
        cols.append((lab + " (%s)" % fmt(mx), mx, "s1"))
    cols.append(("צילום שיעור 1 (12)", 12, "s1tot"))
    for lab, mx in L2S:
        cols.append((lab + " (%s)" % fmt(mx), mx, "s2"))
    cols.append(("צילום שיעור 2 (18)", 18, "s2tot"))
    cols.append(("סה\"כ צילומים (30)", 30, "scr"))
    cols.append(("קובץ צילום זהה בשיעור 1", None, "id1"))
    cols.append(("קובץ צילום זהה בשיעור 2", None, "id2"))
    cols.append(("שם קובץ של אחר", None, "pen"))
    cols.append(("סה\"כ הורדות", None, "ded"))
    cols.append(("ציון סופי (100)", 100, "final"))
    cols.append(("מה ראיתי בצילומים (לעזרה בבדיקה)", None, "obs"))

    # שורת קבוצות
    groups_row = [("", 1, 1)]
    idx = {k: [i + 1 for i, c in enumerate(cols) if c[2] == k] for k in set(c[2] for c in cols)}
    col_of = lambda k: idx[k][0]
    spans = [("חלק כתוב, שיעור 1", 2, col_of("l1")), ("חלק כתוב, שיעור 2", col_of("l1") + 1, col_of("l2")),
             ("סיכום כתוב", col_of("text"), col_of("bonus")),
             ("צילומי מסך (אפשר לשנות)", col_of("s1tot") - 2, col_of("scr")),
             ("הורדות וציון סופי", col_of("id1"), col_of("final"))]
    ws.row_dimensions[1].height = 22
    for lab, a, b in spans:
        ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
        c = ws.cell(1, a, lab)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = H_FILL
        c.alignment = WRAP
    for i, (lab, mx, kind) in enumerate(cols, 1):
        c = ws.cell(2, i, lab)
        c.font = Font(name=FONT, bold=True)
        c.alignment = WRAP
        c.border = BOX
        c.fill = SCR_FILL if kind in ("s1", "s2", "s1tot", "s2tot") else (TOT_FILL if kind in ("l1", "l2", "text", "scr", "final") else TXT_FILL)
    ws.row_dimensions[2].height = 75

    L = lambda k, j=0: get_column_letter(idx[k][j])
    dv = DataValidation(type="list", formula1='"כן,לא"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, st in enumerate(studs, 3):
        pts = [p for p, _ in st["parts"]]
        ws.cell(r, 1, st["name"]).font = Font(name=FONT, bold=True)
        tcols = idx["t"]
        for j, p in enumerate(pts):
            ws.cell(r, tcols[j], p)
        ws.cell(r, col_of("l1"), "=SUM(%s%d:%s%d)" % (get_column_letter(tcols[0]), r, get_column_letter(tcols[5]), r))
        ws.cell(r, col_of("l2"), "=SUM(%s%d:%s%d)" % (get_column_letter(tcols[6]), r, get_column_letter(tcols[11]), r))
        ws.cell(r, col_of("text"), "=%s%d+%s%d" % (L("l1"), r, L("l2"), r))
        b = st.get("bonus", (0, 0, ""))
        ws.cell(r, col_of("bonus"), b[0] + b[1])
        s1a, s1b = L("s1", 0), L("s1", 1)
        s2a, s2b = L("s2", 0), L("s2", 4)
        c1, _, c2, _ = scr(st)
        for j, v in enumerate(c1):
            ws.cell(r, idx["s1"][j], v)
        for j, v in enumerate(c2):
            ws.cell(r, idx["s2"][j], v)
        ws.cell(r, col_of("s1tot"), "=SUM(%s%d:%s%d)" % (s1a, r, s1b, r))
        ws.cell(r, col_of("s2tot"), "=SUM(%s%d:%s%d)" % (s2a, r, s2b, r))
        ws.cell(r, col_of("id1"), "כן" if "L1" in st["identical"] else "לא")
        ws.cell(r, col_of("id2"), "כן" if "L2" in st["identical"] else "לא")
        dv.add("%s%d" % (L("id1"), r))
        dv.add("%s%d" % (L("id2"), r))
        t1, t2, i1, i2 = L("s1tot"), L("s2tot"), L("id1"), L("id2")
        ws.cell(r, col_of("scr"), "=%s%d+%s%d" % (t1, r, t2, r))
        ws.cell(r, col_of("pen"), "כן" if st["foreign"] else "לא")
        dv.add("%s%d" % (L("pen"), r))
        ws.cell(r, col_of("ded"), '=%d*((%s%d="כן")+(%s%d="כן"))+(%s%d="כן")'
                % (G.IDENT_PENALTY, i1, r, i2, r, L("pen"), r))
        ws.cell(r, col_of("final"), "=MAX(0,MIN(100,ROUND(%s%d+%s%d+%s%d-%s%d,0)))"
                % (L("text"), r, L("bonus"), r, L("scr"), r, L("ded"), r))
        ws.cell(r, col_of("obs"), st["obs"])
        for i, (lab, mx, kind) in enumerate(cols, 1):
            c = ws.cell(r, i)
            c.border = BOX
            c.font = Font(name=FONT, bold=(kind in ("name", "text", "final")))
            c.alignment = RIGHT if kind in ("name", "obs") else Alignment(horizontal="center", vertical="center")
            if kind in ("s1", "s2", "s1tot", "s2tot"):
                c.fill = SCR_FILL
            elif kind in ("l1", "l2", "text", "scr", "final", "ded"):
                c.fill = TOT_FILL
        ws.row_dimensions[r].height = 48
    widths = {"name": 16, "t": 10, "l1": 10, "l2": 10, "text": 10, "bonus": 8, "s1": 10, "s1tot": 9, "s2": 10,
              "s2tot": 9, "id1": 10, "id2": 10, "scr": 10, "pen": 9, "ded": 9, "final": 12, "obs": 60}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[c[2]]
    ws.freeze_panes = "B3"
    last = 2 + len(studs)
    n = ws.cell(last + 2, 1, "הנקודות בתאים הצהובים נקבעו לפי המחוון, ואפשר לשנות אותן. הציון הסופי מתעדכן לבד. "
                "כל \"כן\" בעמודות של קובץ זהה ובעמודה \"שם קובץ של אחר\" מוריד נקודה אחת.")
    n.font = Font(name=FONT, italic=True)
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=col_of("final"))
    n.alignment = RIGHT
    ws.row_dimensions[last + 2].height = 36

    # פירוט הערות
    ws2 = wb.create_sheet("פירוט החלק הכתוב")
    ws2.sheet_view.rightToLeft = True
    hdr = ["שם התלמיד", "חלק", "נקודות", "מתוך", "הסבר"]
    for i, h in enumerate(hdr, 1):
        c = ws2.cell(1, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = H_FILL
        c.alignment = WRAP
    r = 2
    for st in studs:
        for (pid, label, mx), (p, key) in zip(G.TEXT_PARTS, st["parts"]):
            for i, v in enumerate([st["name"], label, p, mx, G.N[key]], 1):
                c = ws2.cell(r, i, v)
                c.font = Font(name=FONT)
                c.alignment = RIGHT if i in (1, 2, 5) else Alignment(horizontal="center", vertical="top")
                c.border = BOX
            r += 1
        b = st.get("bonus")
        if b:
            for i, v in enumerate([st["name"], "בונוס", b[0] + b[1], 5, b[2]], 1):
                c = ws2.cell(r, i, v)
                c.font = Font(name=FONT)
                c.alignment = RIGHT
                c.border = BOX
            r += 1
    for i, w in enumerate([16, 34, 8, 7, 90], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    # למורה
    ws3 = wb.create_sheet("למורה")
    ws3.sheet_view.rightToLeft = True
    for i, h in enumerate(["שם התלמיד", "צילום זהה", "שם קובץ של אחר", "הערה למורה (לא לתלמיד)"], 1):
        c = ws3.cell(1, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = H_FILL
        c.alignment = WRAP
    for r, st in enumerate(studs, 2):
        ident = " ו-".join({"L1": "שיעור 1", "L2": "שיעור 2"}[x] for x in st["identical"]) or "לא"
        for i, v in enumerate([st["name"], ident, "כן, נקודה אחת" if st["foreign"] else "לא", st["teacher"]], 1):
            c = ws3.cell(r, i, v)
            c.font = Font(name=FONT)
            c.alignment = RIGHT
            c.border = BOX
    r = len(studs) + 3
    for line in G.CLASS_NOTES:
        c = ws3.cell(r, 1, line)
        c.font = Font(name=FONT, bold=line.endswith(":"))
        ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c.alignment = RIGHT
        ws3.row_dimensions[r].height = 30
        r += 1
    for i, w in enumerate([16, 18, 16, 100], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # מחוון
    ws4 = wb.create_sheet("חלוקת הנקודות")
    ws4.sheet_view.rightToLeft = True
    rows = [("חלק", "נקודות")] + [(l, m) for _, l, m in G.L1_PARTS] + [("שיעור 1: צילום הכרטיס", 12)] + \
           [(l, m) for _, l, m in G.L2_PARTS] + [("שיעור 2: צילום רשימת הצעדים", 18)] + \
           [("סה\"כ חלק כתוב", 70), ("סה\"כ צילומי מסך", 30), ("בונוס", "עד 5"),
            ("קובץ צילום זהה לקובץ של תלמיד אחר", "נקודה אחת לכל שיעור"), ("שם קובץ של תלמיד אחר בצילום", "נקודה אחת מהציון הסופי")]
    for r, (a, b) in enumerate(rows, 1):
        ws4.cell(r, 1, a).font = Font(name=FONT, bold=(r == 1))
        ws4.cell(r, 2, b).font = Font(name=FONT, bold=(r == 1))
        ws4.cell(r, 1).alignment = RIGHT
    ws4.column_dimensions["A"].width = 50
    ws4.column_dimensions["B"].width = 24
    wb.save(path)


# ---------------- Word ----------------
def rtl_par(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(OxmlElement("w:bidi"))
    if p.alignment is None:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def rtl_run(run, size=11, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), FONT)
    rPr.append(OxmlElement("w:rtl"))
    b = OxmlElement("w:szCs")
    b.set(qn("w:val"), str(size * 2))
    rPr.append(b)
    if bold:
        rPr.append(OxmlElement("w:bCs"))


def para(container, text="", size=11, bold=False, color=None, space_after=4, align=None):
    p = container.add_paragraph()
    if align is not None:
        p.alignment = align
    rtl_par(p)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        rtl_run(p.add_run(text), size, bold, color)
    return p


def cell_text(cell, text, bold=False, size=10, fill=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rtl_par(p)
    rtl_run(p.add_run(text), size, bold)
    if fill:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)


def rtl_table(t, widths_cm):
    tblPr = t._tbl.tblPr
    tblPr.append(OxmlElement("w:bidiVisual"))
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)


def build_word(group, studs, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    sectPr = sec._sectPr
    sectPr.append(OxmlElement("w:bidi"))
    st_norm = doc.styles["Normal"]
    st_norm.font.name = FONT
    st_norm.font.size = Pt(11)

    for n, st in enumerate(studs):
        l1, l2, bonus = totals(st)
        text = l1 + l2
        para(doc, "%s, %s" % (TITLE, G.GROUPS[group]), 9, color="666666", space_after=0)
        para(doc, "משוב: %s" % st["name"], 16, True, "1F4E79", space_after=6)
        c1, n1, c2, n2 = scr(st)
        ded_pts = deductions(st)
        para(doc, "ציון סופי: %d מתוך 100" % final_grade(st), 14, True, space_after=2)
        line = "חלק כתוב: %s מתוך 70. צילומי מסך: %s מתוך 30" % (fmt(text), fmt(sum(c1) + sum(c2)))
        if bonus:
            line += ". בונוס: %s" % fmt(bonus)
        if ded_pts:
            line += ". הורדות: %s" % fmt(ded_pts)
        para(doc, line + ".", 10.5, space_after=6)

        rows = [(label, "%s מתוך %s" % (fmt(p), fmt(mx)), G.N[key]) for (pid, label, mx), (p, key) in zip(G.TEXT_PARTS, st["parts"])]
        b = st.get("bonus")
        if b:
            rows.append(("בונוס", "%s מתוך 5" % fmt(b[0] + b[1]), b[2]))
        t = doc.add_table(rows=1 + len(rows), cols=3)
        t.style = "Table Grid"
        for j, h in enumerate(["חלק", "נקודות", "הסבר"]):
            cell_text(t.rows[0].cells[j], h, True, 10, "DDEBF7", center=(j == 1))
        for i, (a, pts, note) in enumerate(rows, 1):
            cell_text(t.rows[i].cells[0], a, False, 9)
            cell_text(t.rows[i].cells[1], pts, True, 9.5, center=True)
            cell_text(t.rows[i].cells[2], note, False, 9)
        rtl_table(t, [4.0, 2.5, 10.9])

        para(doc, "", space_after=2)
        t2 = doc.add_table(rows=4, cols=3)
        t2.style = "Table Grid"
        cell_text(t2.rows[0].cells[0], "צילום מסך, שיעור 1 (הכרטיס)", False, 9, "FFF2CC")
        cell_text(t2.rows[0].cells[1], "%s מתוך 12" % fmt(sum(c1)), True, 9.5, "FFF2CC", center=True)
        cell_text(t2.rows[0].cells[2], n1, False, 9, "FFF2CC")
        cell_text(t2.rows[1].cells[0], "צילום מסך, שיעור 2 (רשימת הצעדים)", False, 9, "FFF2CC")
        cell_text(t2.rows[1].cells[1], "%s מתוך 18" % fmt(sum(c2)), True, 9.5, "FFF2CC", center=True)
        cell_text(t2.rows[1].cells[2], n2, False, 9, "FFF2CC")
        ded = []
        for les in st["identical"]:
            ded.append(G.IDENT_NOTE[les])
        if st["foreign"]:
            ded.append(G.FOREIGN_NOTE)
        cell_text(t2.rows[2].cells[0], "הורדות", False, 9)
        cell_text(t2.rows[2].cells[1], ("%s-" % fmt(ded_pts)) if ded_pts else "0", True, 9.5, center=True)
        cell_text(t2.rows[2].cells[2], " ".join(ded) if ded else "אין", False, 9)
        cell_text(t2.rows[3].cells[0], "ציון סופי", True, 10.5, "E2EFDA")
        cell_text(t2.rows[3].cells[1], "%d מתוך 100" % final_grade(st), True, 10.5, "E2EFDA", center=True)
        cell_text(t2.rows[3].cells[2], "חלק כתוב %s, צילומים %s, בונוס %s, הורדות %s. מעוגל למספר שלם." % (fmt(text), fmt(sum(c1) + sum(c2)), fmt(bonus), fmt(ded_pts)), False, 9, "E2EFDA")
        rtl_table(t2, [4.0, 2.5, 10.9])

        para(doc, "", space_after=2)
        para(doc, "סיכום", 12, True, "1F4E79", space_after=2)
        para(doc, st["summary"], 10.5, space_after=4)
        if n < len(studs) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    scores = {"students": []}
    for g in ("A", "B"):
        studs = [s for s in G.STUDENTS if s["g"] == g]
        name = "%s, %s" % (TITLE, G.GROUPS[g])
        build_excel(g, studs, os.path.join(OUT, "ציונים - %s.xlsx" % name))
        build_word(g, studs, os.path.join(OUT, "משוב לתלמידים - %s.docx" % name))
    for s in G.STUDENTS:
        sc = {pid: p for (pid, _, _), (p, _) in zip(G.TEXT_PARTS, s["parts"])}
        sc.update({"L1.S": 0, "L2.S": 0, "B.L1Q3": s.get("bonus", (0, 0, ""))[0], "B.insight": s.get("bonus", (0, 0, ""))[1]})
        scores["students"].append({"name": s["name"], "folder": G.GROUPS[s["g"]], "scores": sc,
                                   "identical_screens": s["identical"], "foreign_file": s["foreign"]})
        l1, l2, b = totals(s)
        w = words(s["summary"]) + sum(words(G.IDENT_NOTE[x]) for x in s["identical"]) + (words(G.FOREIGN_NOTE) if s["foreign"] else 0)
        c1, _, c2, _ = scr(s)
        print("%-16s text %-5s scr %-5s bonus %s ded %s FINAL %d words %d" % (s["name"], fmt(l1 + l2), fmt(sum(c1) + sum(c2)), fmt(b), deductions(s), final_grade(s), w))
    with open(os.path.join(OUT, "scores.json"), "w", encoding="utf-8") as f:
        json.dump(scores, f, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()
