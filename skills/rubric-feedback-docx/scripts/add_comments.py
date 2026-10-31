"""
add_comments.py: plant grading feedback as editable Word comment bubbles.

Usage:
    python add_comments.py spec.json

spec.json:
{
  "author": "משה זלינגר",
  "initials": "מז",
  "out_dir": "out/תרגיל 3 קבוצה א - עם הערות",     # unzipped folder (the working copy)
  "zip_path": "out/... - עם הערות.zip",             # optional, built from out_dir
  "files": [
    {
      "src": "work/student_x/answer.docx",
      "dst": "student_x/answer.docx",                # relative to out_dir
      "comments": [
        {"anchor": {"type": "first"},                    "head": "ציון: 89 מתוך 100", "body": ["...", "..."]},
        {"anchor": {"type": "p", "para": 10, "marker": "א."}, "head": "שאלה 1א (5 מתוך 6)", "body": "..."},
        {"anchor": {"type": "range", "from": 12, "to": 16}, "head": "...", "body": "..."},
        {"anchor": {"type": "img", "para": 4},          "head": "צילום מסך (15 מתוך 15)", "body": "..."}
      ]
    }
  ]
}

Anchor types:
  first  : first paragraph with text (use for the grade comment)
  p      : whole paragraph; with "marker", from the marker to the end of that line
           (for answers written as several lines inside one paragraph)
  range  : from paragraph "from" to paragraph "to", inclusive
  img    : the first picture in paragraph "para"

Why the files stay editable and savable after manual edits:
  * every comment paragraph gets a w14:paraId, and the package gets
    commentsExtended.xml, commentsIds.xml and people.xml, like a comment typed in Word.
    Without these, some Word versions fail to save edits to comment text.
  * RTL is set on paragraph and run level, in schema order (bidi after pStyle, rtl and
    lang near the end of rPr). Out-of-order elements make Word ignore them or "repair" the file.
  * the student's own text is never rewritten. Runs are only split at anchor boundaries.
  * outputs go to a normal folder. Files opened from inside a ZIP are read-only temp copies,
    which is the most common reason manual edits "do not save".
"""
import copy, json, os, random, sys, zipfile
from datetime import datetime, timezone

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run
from lxml import etree

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
W16CID = "http://schemas.microsoft.com/office/word/2016/wordml/cid"
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
             "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
             "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
             "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang"]


def hexid():
    return "%08X" % random.randint(1, 0x7FFFFFFE)


# ---------- anchoring ----------
def _is_text_run(r):
    return not r._r.xpath(".//w:drawing|.//w:pict|.//w:object|.//w:fldChar|.//w:instrText")


def _split_run(run, off):
    t = run.text
    if 0 < off < len(t):
        new = copy.deepcopy(run._r)
        run._r.addnext(new)
        run.text = t[:off]
        Run(new, run._parent).text = t[off:]


def _span_runs(p, start, end):
    for pos in (start, end):
        acc = 0
        for r in p.runs:
            L = len(r.text)
            if acc < pos < acc + L and _is_text_run(r):
                _split_run(r, pos - acc)
                break
            acc += L
    acc, sel = 0, []
    for r in p.runs:
        L = len(r.text)
        if L and acc >= start and acc + L <= end:
            sel.append(r)
        acc += L
    return sel


def _text_runs(p):
    return [r for r in p.runs if r.text.strip()]


def resolve(doc, a):
    P = doc.paragraphs
    t = a["type"]
    if t == "first":
        for p in P:
            if _text_runs(p):
                return _text_runs(p)
        return P[0].runs[:1]
    if t == "p":
        p = P[a["para"]]
        if a.get("marker"):
            txt = "".join(r.text for r in p.runs)
            s = txt.find(a["marker"])
            if s < 0:
                raise ValueError(f"marker not found: {a}")
            e = txt.find("\n", s)
            return _span_runs(p, s, len(txt) if e < 0 else e)
        return _text_runs(p)
    if t == "range":
        return [_text_runs(P[a["from"]])[0], _text_runs(P[a["to"]])[-1]]
    if t == "img":
        return [[r for r in P[a["para"]].runs if r._r.xpath(".//w:drawing|.//w:pict")][0]]
    raise ValueError(a)


# ---------- RTL comment paragraphs ----------
def _rtl(par, bold=False):
    pPr = par._p.get_or_add_pPr()
    for e in pPr.findall(qn("w:bidi")) + pPr.findall(qn("w:jc")):
        pPr.remove(e)
    bidi = OxmlElement("w:bidi")
    ps = pPr.find(qn("w:pStyle"))
    ps.addnext(bidi) if ps is not None else pPr.insert(0, bidi)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "right")
    rp = pPr.find(qn("w:rPr"))
    rp.addprevious(jc) if rp is not None else pPr.append(jc)
    for r in par.runs:
        rPr = r._r.get_or_add_rPr()
        if bold:
            r.bold = True
            rPr.append(OxmlElement("w:bCs"))
        rPr.append(OxmlElement("w:rtl"))
        lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "he-IL"); rPr.append(lang)
        kids = sorted(rPr, key=lambda e: RPR_ORDER.index(etree.QName(e).localname)
                      if etree.QName(e).localname in RPR_ORDER else 99)
        for k in kids:
            rPr.remove(k)
        for k in kids:
            rPr.append(k)
    par._p.set("{%s}paraId" % W14, hexid())
    par._p.set("{%s}textId" % W14, "77777777")


# ---------- Word 2013+ comment parts ----------
def _add_part(doc, name, ctype, reltype, xml):
    part = Part(PackURI(name), ctype, xml.encode("utf-8"), doc.part.package)
    doc.part.relate_to(part, reltype)


def _modern_parts(doc, author):
    cpart = doc.part._comments_part
    root = cpart.element
    ext, ids = [], []
    for c in root.findall(qn("w:comment")):
        last = c.findall(qn("w:p"))[-1]
        pid = last.get("{%s}paraId" % W14)
        ext.append(f'<w15:commentEx w15:paraId="{pid}" w15:done="0"/>')
        ids.append(f'<w16cid:commentId w16cid:paraId="{pid}" w16cid:durableId="{hexid()}"/>')
    rel = "http://schemas.microsoft.com/office/2011/relationships/"
    _add_part(doc, "/word/commentsExtended.xml",
              "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml",
              rel + "commentsExtended",
              f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w15:commentsEx xmlns:w15="{W15}" '
              f'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">{"".join(ext)}</w15:commentsEx>')
    _add_part(doc, "/word/commentsIds.xml",
              "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsIds+xml",
              "http://schemas.microsoft.com/office/2016/09/relationships/commentsIds",
              f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w16cid:commentsIds xmlns:w16cid="{W16CID}">'
              f'{"".join(ids)}</w16cid:commentsIds>')
    _add_part(doc, "/word/people.xml",
              "application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml",
              rel + "people",
              f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w15:people xmlns:w15="{W15}">'
              f'<w15:person w15:author="{author}"><w15:presenceInfo w15:providerId="None" w15:userId="{author}"/>'
              f'</w15:person></w15:people>')


def comment_file(src, dst, comments, author, initials):
    doc = docx.Document(src)
    if doc.part._comments_part is not None and len(doc.comments) > 0:
        print("  note: file already has comments:", src)
    anchors = [resolve(doc, c["anchor"]) for c in comments]   # resolve all first, indices stay stable
    date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    for c, runs in zip(comments, anchors):
        cm = doc.add_comment(runs, text=c["head"], author=author, initials=initials)
        cm._comment_elm.set(qn("w:date"), date)
        _rtl(cm.paragraphs[0], bold=True)
        body = c.get("body") or []
        for line in ([body] if isinstance(body, str) else body):
            _rtl(cm.add_paragraph(line))
    _modern_parts(doc, author)
    os.makedirs(os.path.dirname(dst) or ".", exist_ok=True)
    doc.save(dst)


def main(spec_path):
    spec = json.load(open(spec_path, encoding="utf-8"))
    out = spec["out_dir"]
    for f in spec["files"]:
        dst = os.path.join(out, f["dst"])
        comment_file(f["src"], dst, f["comments"], spec["author"], spec.get("initials", ""))
        docx.Document(dst)          # reopen check
        print("ok", f["dst"], len(f["comments"]))
    if spec.get("zip_path"):
        with zipfile.ZipFile(spec["zip_path"], "w", zipfile.ZIP_DEFLATED) as z:
            for root, _, files in os.walk(out):
                for n in files:
                    if n.endswith(".docx"):
                        p = os.path.join(root, n)
                        z.write(p, os.path.relpath(p, out))
        print("zip:", spec["zip_path"])


if __name__ == "__main__":
    main(sys.argv[1])
